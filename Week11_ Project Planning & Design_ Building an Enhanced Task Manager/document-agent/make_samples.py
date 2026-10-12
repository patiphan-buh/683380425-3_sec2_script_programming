# Creates sample input files: PDFs in documents/ and templates/report_template.docx
import os
import docx
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

BASE = os.path.dirname(os.path.abspath(__file__))
D = os.path.join(BASE, 'documents')

def pdf(name, pages):
    c = canvas.Canvas(os.path.join(D, name), pagesize=A4)
    for lines in pages:
        y = 780
        for ln in lines:
            c.setFont("Helvetica-Bold" if ln.isupper() or ln.endswith(":") else "Helvetica", 11)
            c.drawString(60, y, ln); y -= 18
        c.showPage()
    c.save()

pdf('dummy_doc1.pdf', [["This is Dummy Document 1 for merging test."]])
pdf('dummy_doc2.pdf', [["This is Dummy Document 2 for merging test."]])
pdf('source_document.pdf', [[
    "QUARTERLY PROCESS REVIEW", "",
    "Purpose:",
    "This document is the source file used to demonstrate PDF text extraction",
    "and watermarking in the document automation agent.", "",
    "Highlights:",
    "- Manual document handling was reduced by automating repetitive tasks.",
    "- Reports are generated from a Word template with placeholders.",
    "- Every run of the agent writes an audit report for traceability.", "",
    "Conclusion:",
    "Automating document workflows saves time and keeps output consistent."]])

t = docx.Document()
t.add_heading('{{title}}', level=1)
t.add_paragraph('Author: {{author}}')
t.add_paragraph('Date: {{date}}')
t.add_heading('Summary', level=2)
t.add_paragraph('{{summary}}')
tbl = t.add_table(rows=3, cols=2); tbl.style = 'Table Grid'
for r, (a, b) in enumerate([("Field", "Value"), ("Author", "{{author}}"), ("Generated", "{{date}}")]):
    tbl.rows[r].cells[0].text = a; tbl.rows[r].cells[1].text = b
os.makedirs(os.path.join(BASE, 'templates'), exist_ok=True)
t.save(os.path.join(BASE, 'templates', 'report_template.docx'))
