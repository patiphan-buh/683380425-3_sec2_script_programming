import datetime
import docx
from utils import setup_logging

logger = setup_logging(__name__)


def _fill_paragraph(paragraph, values):
    for key, val in values.items():
        token = "{{" + key + "}}"
        if token in paragraph.text:
            new_text = paragraph.text.replace(token, str(val))
            for run in paragraph.runs[1:]:
                run.text = ""
            if paragraph.runs:
                paragraph.runs[0].text = new_text
            else:
                paragraph.add_run(new_text)


def generate_report(template_path, output_path, data):
    """Fill {{placeholders}} in a Word template and save the result."""
    document = docx.Document(template_path)
    values = {"date": datetime.date.today().isoformat(), **data}
    for p in document.paragraphs:
        _fill_paragraph(p, values)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    _fill_paragraph(p, values)
    document.save(output_path)
    logger.info(f"Generated report from template -> {output_path}")


def extract_word_text(docx_path, output_txt):
    document = docx.Document(docx_path)
    lines = [p.text for p in document.paragraphs if p.text.strip()]
    for t in document.tables:
        for row in t.rows:
            lines.append(" | ".join(c.text for c in row.cells))
    with open(output_txt, 'w', encoding='utf-8') as f:
        f.write("\n".join(lines))
    logger.info(f"Extracted {len(lines)} lines -> {output_txt}")
    return len(lines)
